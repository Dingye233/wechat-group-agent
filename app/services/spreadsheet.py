"""Safe spreadsheet generation for group tasks.

This module only writes inside ``data/agent_workspace/exports``.  The group
Agent never gets arbitrary filesystem access; it can ask for a table to be
built from structured rows and the runtime sends the resulting file through
the normal wechatauto-replica ``SendFiles`` path.
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[2]
EXPORT_ROOT = PROJECT_ROOT / "data" / "agent_workspace" / "exports"

MAX_SHEETS = 10
MAX_ROWS = 5000
MAX_COLUMNS = 50
MAX_FILENAME = 80
_INVALID_FILENAME = re.compile(r'[\\/:*?"<>|\r\n]+')
_INVALID_SHEET = re.compile(r"[\[\]:*?/\\]+")


def _safe_filename(value: Any) -> str:
    name = Path(str(value or "")).name
    name = _INVALID_FILENAME.sub("_", name).strip().strip(".")
    if not name:
        name = f"表格_{datetime.now():%Y%m%d_%H%M%S}"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    stem = Path(name).stem[:MAX_FILENAME] or "表格"
    return f"{stem}.xlsx"


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, datetime):
        return value
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def _normalize_sheets(args: dict[str, Any]) -> list[dict[str, Any]]:
    raw = args.get("sheets")
    if not isinstance(raw, list) or not raw:
        rows = args.get("rows")
        if isinstance(rows, list) and rows:
            raw = [{"name": args.get("sheet_name") or args.get("name") or "Sheet1", "rows": rows}]
        else:
            raise ValueError("缺少 sheets：请传 [{name, rows}]，rows 是二维数组")
    sheets: list[dict[str, Any]] = []
    for index, item in enumerate(raw[:MAX_SHEETS]):
        if not isinstance(item, dict):
            raise ValueError(f"sheets[{index}] 必须是对象")
        rows = item.get("rows")
        if not isinstance(rows, list):
            raise ValueError(f"sheets[{index}].rows 必须是二维数组")
        name = _INVALID_SHEET.sub("_", str(item.get("name") or f"Sheet{index + 1}"))
        name = (name.strip() or f"Sheet{index + 1}")[:31]
        sheets.append(
            {
                "name": name,
                "rows": rows[:MAX_ROWS],
                "header": bool(item.get("header", True)),
            }
        )
    if not sheets:
        raise ValueError("没有可写入的工作表")
    return sheets


def create_excel(args: dict[str, Any] | None = None) -> dict[str, Any]:
    """Create an .xlsx file from rows and return its local path."""

    args = args or {}
    sheets = _normalize_sheets(args)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise RuntimeError("缺少 openpyxl，请先安装：pip install openpyxl") from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    total_rows = 0
    total_cells = 0

    for sheet in sheets:
        worksheet = workbook.create_sheet(title=sheet["name"])
        rows = sheet["rows"]
        header = sheet["header"]
        max_columns_used = 0
        for row_index, row in enumerate(rows, start=1):
            if isinstance(row, dict):
                values = list(row.values())
            elif isinstance(row, (list, tuple)):
                values = list(row)
            else:
                values = [row]
            values = values[:MAX_COLUMNS]
            max_columns_used = max(max_columns_used, len(values))
            total_cells += len(values)
            if total_cells > MAX_ROWS * MAX_COLUMNS:
                raise ValueError("表格内容过大，请拆成多个文件")
            for column_index, value in enumerate(values, start=1):
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=_cell_value(value),
                )
        total_rows += len(rows)

        if rows and header:
            for cell in worksheet[1]:
                cell.font = Font(bold=True, name="Microsoft YaHei")
                cell.fill = PatternFill("solid", fgColor="DDEBF7")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            worksheet.freeze_panes = "A2"
            try:
                worksheet.auto_filter.ref = worksheet.dimensions
            except Exception:
                pass

        # Keep the sheet readable without letting one huge cell dominate.
        for column_index in range(1, max_columns_used + 1):
            letter = get_column_letter(column_index)
            longest = 0
            for row_index in range(1, min(len(rows), MAX_ROWS) + 1):
                value = worksheet.cell(row=row_index, column=column_index).value
                longest = max(longest, len(str(value or "")))
            worksheet.column_dimensions[letter].width = min(max(9, longest * 1.7), 55)

        if rows:
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

    if not workbook.worksheets:
        raise ValueError("没有可写入的表格数据")

    EXPORT_ROOT.mkdir(parents=True, exist_ok=True)
    filename = _safe_filename(args.get("filename"))
    target = (EXPORT_ROOT / filename).resolve()
    if EXPORT_ROOT.resolve() not in target.parents:
        raise ValueError("文件名不安全")
    if target.exists():
        target = target.with_name(
            f"{target.stem}_{datetime.now():%H%M%S}{target.suffix}"
        )
    workbook.save(target)
    return {
        "path": str(target),
        "filename": target.name,
        "sheets": len(workbook.worksheets),
        "rows": total_rows,
        "bytes": target.stat().st_size,
        "message": "Excel 已生成，发送时会自动附到群聊消息后面",
    }
